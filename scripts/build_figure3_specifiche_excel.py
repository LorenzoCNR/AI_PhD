from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results' / 'figura_3_2026-10-28' / 'Figura_3_specifiche_modelli.xlsx'
HEAD = ['Modello / sottoplot', 'Dati e split del fit', 'Dati visualizzati / test', 'Rete e input', 'Dimensione latente fit', 'Run / seed', 'Parametri rete', 'Iperparametri', 'Loss / similarità', 'Label nel fit', 'Batch size', 'Coordinate / decoding', 'Paper e limiti', 'Fonti locali']
ALLA='193 trial attivi completi; 600 bin/trial; 65 neuroni; 1 ms/bin'
ALLP='171 trial passivi completi; 600 bin/trial; 65 neuroni; 1 ms/bin'
ALLB='364 trial attivi + passivi completi; 600 bin/trial; 65 neuroni; 1 ms/bin'
TR='Train 141 trial attivi; valid 26; test 26, split per trial'
MIX='Train 272 trial attivi + passivi; valid 46; test 46, split per trial'
CF='cebra_replica/monkey/fit_figure3_paper.py; results/monkey_figure3_paper_v1/stage01_fit'
HF='cebra_replica/monkey/fit_figure3h.py; results/monkey_figure3_paper_v1/stage01_fit_heldout'
SF='results/user_soft_monkey_v1/PROTOCOL.md; stage01_fit/observed; stage02_embedding/observed'
CP='LR 3e-4; T=1; cosine; time_offsets=10; time_delta per label, time per temporale'
SP='Adam LR 3e-4; beta 0,9/0,999; eps 1e-8; wd 0; T embedding 0,1; T metadata 0,5; no early stopping'
V='Plot descrittivo: righe valide train + validation + test; encoder fit solo su train'
PAPER='Dimensioni/etichette del paper come da caption; gli iperparametri originali non sono tutti confermati'
rows={x:[] for x in 'bcdefghi'}
def add(p,*v):
    assert len(v)==14,(p,len(v))
    rows[p].append(list(v))
def c(p,name,data,label,dim,plot,note=PAPER,src=CF):
    add(p,name,data,'Intero supporto del fit; descrittivo','CEBRA offset10; input 10 bin × 65 neuroni',dim,'10.000 update; seed 42','Hidden: non esplicitato nel manifest',CP,'InfoNCE; cosine',label,512,plot,note,src)
def s(p,arch,name,label,dim,plot,note):
    desc='CNN offset10, hidden 32' if arch=='CNN' else 'Transformer width 64, 4 teste, 2 layer, dropout 0,1'
    support='Test held-out: 26 trial attivi; decoder addestrato sul train' if p in ('h','i') else V
    add(p,'Soft '+arch+' — '+name,TR,support,desc+'; input 10 bin × 65 neuroni',dim,'10.000 update; seed 42–46',desc,SP,'Stable Soft; embedding normalizzato L2',label,512,plot,note,SF)
c('b','CEBRA Direction',ALLA,'Direzione: 8 categorie','4D','Media per direzione; z3, z1')
c('b','CEBRA Time',ALLA,'Tempo; nessuna label comportamentale','4D','Media per direzione; z1, z2')
for lbl in ('con label in inferenza','senza label in inferenza'):
    add('b','conv-piVAE '+lbl,TR,'Plot trial attivi; inferenza dal medesimo checkpoint','CNN piVAE; finestre 10 × 65','4D','600 epoche; seed 42; checkpoint migliore epoca 584','Prior discreto 8 categorie; gen_nodes 60; 2 flow block','Adam LR 5e-4; batch train 423 × 200; valid 78 × 200','VAE; likelihood Poisson','Direzione nel prior; '+lbl,200,'Media per direzione; z3, z4','Stesso training in entrambe le righe; i parametri originali del paper non sono tutti verificati','cebra_replica/monkey/fit_figure3b_pivae.py; stage01_fit_baselines/conv_pivae/seed_42/FIT.json')
add('b','t-SNE',ALLA,'Interi trial attivi; media per direzione','openTSNE; input 65 feature per bin','2D','1000 iterazioni; seed 42','Non applicabile','cosine; perplexity 30; PCA init; LR auto','Non applicabile','Nessuna',None,'z1, z2','Parametri locali trasferiti dal ratto, non verificati sul paper monkey','cebra_replica/monkey/fit_figure3b_baselines.py')
add('b','UMAP',ALLA,'Interi trial attivi; media per direzione','UMAP; input 65 feature per bin','2D','Seed 42; epoche non specificate','Non applicabile','cosine; n_neighbors 24; min_dist 0,0001','Non applicabile','Nessuna',None,'z1, z2','Parametri locali trasferiti dal ratto, non verificati sul paper monkey','cebra_replica/monkey/fit_figure3b_baselines.py')
add('b','autoLFADS', 'Non disponibile localmente','Non disponibile','autoLFADS','Paper 4D','Non disponibile','Non disponibile','Non disponibile','Non disponibile','Nessuna label comportamentale',None,'Paper: due coordinate','NON STIMATO localmente; non attribuire iperparametri','CEBRA paper Figure 3b')
for arch in ('CNN','Transformer'):
    for n,l in [('Time','Tempo, peso 1'),('Position X,Y','Posizione X,Y, peso 1'),('Time + Position X,Y','Tempo + posizione X,Y; pesi 0,5 e 0,5')]:
        s('b',arch,n,l,'4D','Media direzione; z1, z2','Campagna precedente senza Direction; fit Direction distinto nella riga dedicata')
    desc='CNN offset10, hidden 32' if arch=='CNN' else 'Transformer width 64, 4 teste, 2 layer, dropout 0,1'
    state='Fit completato; dati interamente attivi, senza test held-out' if (ROOT/'results/user_soft_monkey_figure3b_direction_4d_v1'/arch.lower()/'seed_42'/'manifest.json').is_file() else 'Fit in corso; verificare il manifest prima di considerarlo completato'
    add('b','Soft '+arch+' — Direction (nuovo)',ALLA,'Tutti i trial attivi; descrittivo',desc+'; 10 bin × 65','4D','10.000 update; seed 42',desc,'Adam LR 3e-4; T embedding 1; T metadata 0,5','Stable Soft; similarità categoriale','Direzione, 8 categorie',512,'Traiettorie medie per direzione; z1, z2',state,'user_soft/monkey_figure3b_direction_full.py; results/user_soft_monkey_figure3b_direction_4d_v1')
c('c','CEBRA Position X,Y',ALLA,'Posizione continua X,Y','3D','z1, z2, z3; colore posizione X oppure Y','Locale 3D; dimensione nativa originale paper non verificata')
c('d','CEBRA Time',ALLA,'Tempo; X,Y solo per il colore','3D','z1, z2, z3; colore posizione X oppure Y','Locale 3D; dimensione nativa originale paper non verificata')
c('e','CEBRA Direction — attivo',ALLA,'Direzione: 8 categorie','4D','Tre coordinate; media per direzione; attivo')
c('e','CEBRA Direction — passivo',ALLP,'Direzione: 8 categorie','4D','Tre coordinate; media per direzione; passivo')
c('f','CEBRA Direction + Active/Passive',ALLB,'Direzione × condizione: 16 classi','4D','Tre coordinate; grafici attivo e passivo, stesso fit')
c('g','CEBRA Position X,Y — fit corretto',ALLB,'Solo posizione continua X,Y','4D','Tre coordinate; grafici attivo e passivo','Fit corretto position-only; vecchio fit position + active/passive non conforme','results/monkey_figure3_paper_v1/figure3g_position_only_correction_v1/stage01_fit/figure3g_all_position_only_4d_corrected/seed_42/FIT_PROVENANCE.json')
for arch in ('CNN','Transformer'):
    s('c',arch,'Position X,Y','Posizione X,Y, peso 1','3D','z1, z2, z3; colore posizione X o Y','Supporto di fit diverso dal CEBRA locale')
    s('d',arch,'Time','Tempo, peso 1','3D','z1, z2, z3; colore posizione dopo il fit','Supporto di fit diverso dal CEBRA locale')
    s('e',arch,'Position X,Y — solo attivo','Posizione X,Y, peso 1','3D','Tre coordinate; direzione usata dopo il fit','NON EQUIVALENTE: paper richiede direction 4D e fit attivo e passivo indipendenti')
    s('f',arch,'Time + Position X,Y — solo attivo','Tempo + posizione X,Y; pesi 0,5 e 0,5','3D','Tre coordinate; direzione usata dopo il fit','NON EQUIVALENTE: mancano passivi e direction + active/passive')
    s('g',arch,'Position X,Y — solo attivo','Posizione X,Y, peso 1','3D','Tre coordinate; direzione usata dopo il fit','NON EQUIVALENTE: paper usa attivi + passivi e 4D')
for name,data,label,task in [('Position',TR,'Posizione X,Y','R² posizione XY'),('Direction',TR,'Direzione, 8 categorie','Accuratezza direzione'),('Active/Passive',MIX,'Posizione X,Y nel nostro fit locale','Accuratezza attivo/passivo')]:
    add('h','CEBRA — '+name,data,'Test held-out: 26 attivi o 46 complessivi', 'CEBRA offset10; 10 bin × 65', '3, 4, 8, 16, 32D', '10.000 update; 5 seed 42–46 per dimensione','Hidden non esplicitato nel manifest',CP,'InfoNCE; cosine',label,512,task+'; kNN coseno k=36 fit su train','k=36 locale; esatto k originale non verificato. AP: label fit originale non confermata',HF)
for arch in ('CNN','Transformer'):
    for name,label in [('Time','Tempo, peso 1'),('Position X,Y','Posizione X,Y, peso 1'),('Time + Position X,Y','Tempo + posizione X,Y; pesi 0,5 e 0,5')]:
        s('h',arch,name,label,'3, 4, 8, 16, 32D','Metriche held-out su 26 trial attivi test; kNN 36 coseno','Solo attivi: non è confronto active/passive')
add('i','CEBRA Position X,Y',TR,'26 trial attivi test; esempio e score salvati','CEBRA offset10; 10 bin × 65','16D','10.000 update; seed 42','Hidden non esplicitato nel manifest',CP,'InfoNCE; cosine','Posizione X,Y',512,'Traiettoria reale XY e previsione XY; test held-out','Paper: modello 16D; R² locale e paper distinti',HF+'; results/PAPER_FIGURE_DELIVERY/Figure_3/FIGURE3I_RESULT.json')
for arch in ('CNN','Transformer'):
    for name,label in [('Time','Tempo, peso 1'),('Position X,Y','Posizione X,Y, peso 1'),('Time + Position X,Y','Tempo + posizione X,Y; pesi 0,5 e 0,5')]:
        s('i',arch,name,label,'3D nei plot esistenti','Traiettoria XY reale e predetta su test; kNN 36','Fit 3D, diverso dal 16D del paper')

book=Workbook(); sh=book.active; sh.title='Figura 3'; sh.sheet_view.showGridLines=False; sh.freeze_panes='B5'
sh.append(['FIGURA 3 — SPECIFICHE DEI MODELLI LOCALI'])
sh.append(['Dati: scimmia Area 2; 65 canali; 600 bin/trial; 1 ms/bin. I valori locali non verificano automaticamente le configurazioni originali del paper.'])
sh.append(['Split attivi: 141 train / 26 validation / 26 test; tutti i trial: 272 / 46 / 46. Soft storici: T embedding 0,1; T metadata 0,5.'])
widths=[25,26,25,25,15,19,22,31,22,25,12,29,32,18]
for i,w in enumerate(widths,1): sh.column_dimensions[get_column_letter(i)].width=w
for cell in sh[1]: cell.font=Font(name='Arial',size=14,bold=True,color='143050')
for row in (2,3): sh.cell(row,1).font=Font(name='Arial',size=10,italic=True,color='43556B'); sh.row_dimensions[row].height=20
for panel, data in rows.items():
    sh.append(['']); sh.append(['PANNELLO 3'+panel.upper()]); sec=sh.max_row
    for cell in sh[sec][:14]: cell.fill=PatternFill('solid',fgColor='E6E6E6'); cell.font=Font(name='Arial',size=11,bold=True,color='222222')
    sh.row_dimensions[sec].height=22
    sh.append(HEAD); header=sh.max_row
    for cell in sh[header]: cell.fill=PatternFill('solid',fgColor='D0D0D0'); cell.font=Font(name='Arial',size=10,bold=True,color='222222'); cell.alignment=Alignment(wrap_text=True,vertical='center')
    sh.row_dimensions[header].height=30
    for idx,record in enumerate(data):
        sh.append(record); n=sh.max_row
        for cell in sh[n]:
            cell.font=Font(name='Arial',size=10,color='1D2935'); cell.alignment=Alignment(wrap_text=True,vertical='top')
        source=sh.cell(n,14)
        if source.value:
            source.comment=Comment(str(source.value),'Fonti')
            source.value='Fonti: vedi nota'
        if 'NON EQUIVALENTE' in str(record[12]) or 'NON STIMATO' in str(record[12]): sh.cell(n,13).font=Font(name='Arial',size=10,bold=True,color='A2352A')
        import math
        line_count=max(sum(max(1,math.ceil(len(piece)/max(widths[i]-3,1))) for piece in str(value or '').split('\n')) for i,value in enumerate(record[:13]))
        sh.row_dimensions[n].height=max(30,min(90,15*line_count+3))
OUT.parent.mkdir(parents=True,exist_ok=True); book.save(OUT)
check=load_workbook(OUT,read_only=True,data_only=True)
assert check.sheetnames==['Figura 3']
assert sum(1 for r in check.active.iter_rows(values_only=True) if isinstance(r[0],str) and r[0].startswith('PANNELLO 3'))==8
print('SAVED',OUT,'ROWS',check.active.max_row,'MODELS',sum(map(len,rows.values())))
