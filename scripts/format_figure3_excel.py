from pathlib import Path
from copy import copy
import json
import textwrap
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / 'results/figura_3_2026-10-28/Figura_3_specifiche_modelli.xlsx'
FITS = ROOT / 'results/user_soft_monkey_figure3b_direction_4d_v1'
book = load_workbook(PATH)
assert book.sheetnames == ['Figura 3']
sheet = book.active
sheet.sheet_view.showGridLines = True
widths = [29, 29, 28, 28, 16, 23, 25, 38, 25, 30, 13, 32, 35, 19]
for col, width in enumerate(widths, 1):
    sheet.column_dimensions[chr(64 + col)].width = width

line = Side(style='thin', color='888888')
border = Border(left=line, right=line, top=line, bottom=line)
section = header = False
index = 0
for row in sheet.iter_rows(min_row=4, max_row=sheet.max_row, min_col=1, max_col=14):
    title = row[0].value
    if isinstance(title, str) and title.startswith('PANNELLO 3'):
        section, header, index = True, False, 0
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='D3D3D3')
            cell.border = border
        continue
    if section and title == 'Modello / sottoplot':
        header = True
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='BFBFBF')
            cell.font = Font(name='Arial', size=10, bold=True, color='111111')
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 32
        continue
    if title is None:
        section = header = False
        continue
    if not header:
        continue
    shade = 'FFFFFF' if index % 2 == 0 else 'EAF1F7'
    for cell in row:
        cell.border = border
        cell.fill = PatternFill('solid', fgColor=shade)
        cell.alignment = Alignment(vertical='top', wrap_text=True)
    for architecture, needle in [('cnn','Soft CNN — Direction (nuovo)'), ('transformer','Soft Transformer — Direction (nuovo)')]:
        if title == needle:
            manifest = json.loads((FITS / architecture / 'seed_42/manifest.json').read_text(encoding='utf-8'))
            cfg = manifest['config']
            assert manifest['actual_updates'] == 10000 and cfg['output_dimension'] == 4
            assert manifest['training_rows'] == 115800 and cfg['metadata_specs'][0]['geometry'] == 'categorical'
            row[1].value = '193 trial attivi completi; 115.800 bin; 65 canali; 1 ms/bin'
            row[2].value = 'Tutti i trial attivi; fit descrittivo, non held-out'
            row[4].value = '4D nativo'
            row[5].value = f"{manifest['actual_updates']:,} update; seed {cfg['seed']}".replace(',', '.')
            row[7].value = 'Adam LR 3e-4; beta 0,9/0,999; eps 1e-8; weight decay 0; T embedding 1; T metadata 0,5; no early stopping'
            row[8].value = 'Stable Soft; similarità coseno tra embedding; direzione categoriale'
            row[9].value = 'Direzione: 8 categorie; peso 1'
            row[10].value = cfg['batch_size']
            row[11].value = 'Embedding salvato; plot medio z1/z2 non ancora generato'
            row[12].value = 'Fit completato; 193 attivi; nessun test held-out. Paper 3b: CEBRA Direction 4D.'
    lines = 1
    for cell, width in zip(row, widths):
        if cell.value is not None:
            for paragraph in str(cell.value).splitlines() or ['']:
                lines = max(lines, len(textwrap.wrap(paragraph, width=max(8, width - 5), break_long_words=True)) or 1)
    sheet.row_dimensions[row[0].row].height = max(30, 13 * lines + 6)
    index += 1

book.save(PATH)
check = load_workbook(PATH)
s = check.active
assert s.sheet_view.showGridLines
found = [r[0].value for r in s.iter_rows(min_row=1) if 'Direction (nuovo)' in str(r[0].value)]
assert len(found) == 2
assert all(s.cell(r, 1).border.left.style == 'thin' for r in (7, 8))
print('VERIFIED', PATH, 'direction_rows', found, 'sheet_rows', s.max_row)
